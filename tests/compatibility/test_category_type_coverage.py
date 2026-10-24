"""分类类型覆盖：真实API配合内存库，不读写任何运行中的服务。"""

import csv
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from test_baishi_aa import client as _client
from test_baishi_aa import upload, workbook
from test_baishi_review import confirm_execute
from test_import_csv import _iter_sse, _login, _make_ledger

from src.database import get_db
from src.main import app
from src.models import Ledger, ReadTxProjection, SyncChange, User
from src.services.import_data.category_repair import apply_repair, preview_repair

client = _client


def write(client, headers, ledger, resource, **body):
    """通过正式写入接口生成同步与投影，不手工拼造计数。"""
    db = next(app.dependency_overrides[get_db]())
    try:
        body['base_change_id'] = db.scalar(select(func.max(SyncChange.change_id))) or 0
    finally:
        db.close()
    result = client.post(f'/api/v1/write/ledgers/{ledger}/{resource}', headers=headers, json=body)
    assert result.status_code == 200, result.text
    return result.json()['entity_id']


def test_transfer_counts_filter_export_gallery_and_scope(client):
    token = _login(client)
    h = {'Authorization': f'Bearer {token}'}
    first, second = _make_ledger(client, token), _make_ledger(client, token)
    sid = write(client, h, first, 'categories', name='合成转账', kind='transfer')
    assert next(c for c in client.get('/api/v1/read/workspace/categories', headers=h).json()
                if c['id'] == sid)['tx_count'] == 0
    expense = write(client, h, first, 'categories', name='合成支出', kind='expense')
    # 有ID、无ID转账必须各计一次；普通支出不能混入。
    for ledger, category_id in [(first, None), (first, sid), (second, None)]:
        write(client, h, ledger, 'transactions', tx_type='transfer', amount=10,
              happened_at='2026-01-02T10:00:00Z', category_id=category_id,
              attachments=[{'fileName': f'{ledger}-{category_id}.jpg'}])
    write(client, h, first, 'transactions', tx_type='expense', amount=5,
          happened_at='2026-01-02T10:00:00Z', category_id=expense)
    for ledger, expected in [(None, 3), (first, 2), (second, 1)]:
        params = {'ledger_id': ledger} if ledger else {}
        cats = client.get('/api/v1/read/workspace/categories', headers=h, params=params).json()
        assert next(c for c in cats if c['id'] == sid)['tx_count'] == expected
        params['category_sync_id'] = sid
        page = client.get('/api/v1/read/workspace/transactions', headers=h, params=params).json()
        assert page['total'] == expected
        assert all(t['tx_type'] == 'transfer' for t in page['items'])
        exported = client.get('/api/v1/read/workspace/transactions.csv', headers=h, params=params)
        assert exported.status_code == 200, exported.text
        assert len(list(csv.reader(io.StringIO(exported.content.decode('utf-8-sig'))))) == expected + 1
    gallery = client.get('/api/v1/read/workspace/attachments', headers=h,
                         params={'ledger_id': first, 'category_sync_id': sid}).json()
    assert gallery['transaction_total'] == 2
    assert gallery['total'] == 2
    empty = client.get('/api/v1/read/workspace/transactions', headers=h,
                       params={'category_sync_id': sid, 'tx_type': 'income'}).json()
    assert empty['total'] == 0
    page = client.get('/api/v1/read/workspace/transactions', headers=h,
                      params={'category_sync_id': sid, 'limit': 1, 'offset': 1}).json()
    assert page['total'] == 3 and len(page['items']) == 1
    assert client.get('/api/v1/read/workspace/transactions', headers=h,
                      params={'category_sync_id': sid,
                              'date_from': '2026-01-03T00:00:00Z'}).json()['total'] == 0
    # 普通分类仍为精确ID匹配；无分类的收入不能被转账筛选吸收。
    assert client.get('/api/v1/read/workspace/transactions', headers=h,
                      params={'category_sync_id': expense}).json()['total'] == 1
    other = _login(client, 'other-category@example.com')
    other_h = {'Authorization': f'Bearer {other}'}
    own = _make_ledger(client, other)
    other_sid = write(client, other_h, own, 'categories', name='其它用户转账', kind='transfer')
    write(client, other_h, own, 'transactions', tx_type='transfer', amount=20,
          happened_at='2026-01-02T10:00:00Z')
    assert client.get('/api/v1/read/workspace/transactions', headers=other_h,
                      params={'category_sync_id': sid}).json()['total'] == 0
    assert client.get('/api/v1/read/workspace/transactions', headers=h,
                      params={'category_sync_id': sid}).json()['total'] == 3
    # 现有管理员只管理服务，不能通过user_id参数查看其它用户账单。
    db = next(app.dependency_overrides[get_db]())
    try:
        db.scalar(select(User).where(User.email == 'imp@test.com')).is_admin = True
        other_id = db.scalar(select(User.id).where(User.email == 'other-category@example.com'))
        db.commit()
    finally:
        db.close()
    params = {'user_id': other_id, 'category_sync_id': other_sid}
    assert client.get('/api/v1/read/workspace/transactions', headers=h, params=params).json()['total'] == 0
    cats = client.get('/api/v1/read/workspace/categories', headers=h,
                      params={'user_id': other_id}).json()
    assert all(c['id'] != other_sid for c in cats)


def test_csv_transfer_without_category_id_is_readable(client):
    token = _login(client)
    h = {'Authorization': f'Bearer {token}'}
    ledger = _make_ledger(client, token)
    sid = write(client, h, ledger, 'categories', name='合成转账', kind='transfer')
    payload = ('Type,Category,Subcategory,Amount,Account,From Account,To Account,Note,Time,Tags,Attachments\n'
               'Transfer,,,20,,合成A,合成B,,2026-01-02 10:30:00,,\n').encode()
    uploaded = client.post('/api/v1/import/upload', headers=h,
                           files={'file': ('transfer.csv', payload, 'text/csv')},
                           data={'target_ledger_id': ledger})
    assert uploaded.status_code == 200, uploaded.text
    result = client.post(f"/api/v1/import/{uploaded.json()['import_token']}/execute", headers=h)
    assert any(e['event'] == 'complete' for e in _iter_sse(result))
    rows = client.get('/api/v1/read/workspace/transactions', headers=h,
                      params={'category_sync_id': sid}).json()
    assert rows['total'] == 1 and rows['items'][0]['tx_type'] == 'transfer'
    db = next(app.dependency_overrides[get_db]())
    try:
        tx = db.scalar(select(ReadTxProjection))
        assert tx.category_sync_id is None
        assert not db.get(SyncChange, tx.source_change_id).payload_json.get('categoryId')
    finally:
        db.close()


@pytest.mark.parametrize('source', ['csv', 'baishi'])
@pytest.mark.parametrize('legacy', [False, True])
def test_income_import_and_legacy_repair(client, monkeypatch, source, legacy):
    token = _login(client)
    h = {'Authorization': f'Bearer {token}'}
    external = _make_ledger(client, token)
    # 同名收支分类：必须按kind隔离，不能把收入挂到支出。
    write(client, h, external, 'categories', name='合成奖金', kind='expense')
    if source == 'csv':
        payload = ('Type,Category,Subcategory,Amount,Account,From Account,To Account,Note,Time,Tags,Attachments\n'
                   'Income,合成奖金,,12.30,,,,,2026-01-02 10:30:00,,\n').encode()
        result = client.post('/api/v1/import/upload', headers=h,
                             files={'file': ('income.csv', payload, 'text/csv')},
                             data={'target_ledger_id': external})
        assert result.status_code == 200, result.text
        uploaded = result.json()
    else:
        book = load_workbook(io.BytesIO(workbook()))
        sheet = book.active
        sheet.cell(3, 1, '收入')
        sheet.cell(3, 12, '12.30')
        sheet.cell(3, 4, '合成奖金')
        sheet.cell(3, 5).value = None
        output = io.BytesIO()
        book.save(output)
        uploaded = upload(client, token, output.getvalue(), external)
        assert not uploaded.get('parse_errors'), uploaded
    with monkeypatch.context() as patch:
        if legacy:
            patch.setattr('src.routers.import_data.endpoints.resolve_category_id', lambda *a, **k: None)
        if source == 'baishi':
            review = client.post(f"/api/v1/import/{uploaded['import_token']}/review", headers=h, json={})
            assert review.status_code == 200, (review.text, uploaded)
            confirm_execute(client, token, uploaded)
        else:
            response = client.post(f"/api/v1/import/{uploaded['import_token']}/execute", headers=h)
            assert any(e['event'] == 'complete' for e in _iter_sse(response))
    db = next(app.dependency_overrides[get_db]())
    try:
        ledger = db.scalar(select(Ledger).where(Ledger.external_id == external))
        tx = db.scalar(select(ReadTxProjection).where(ReadTxProjection.ledger_id == ledger.id))
        original = dict(db.get(SyncChange, tx.source_change_id).payload_json)
        if legacy:
            assert tx.category_sync_id is None
            plan = preview_repair(db, ledger)
            assert plan['repair_count'] == 1
            apply_repair(db, ledger, plan['plan_hash'])
            db.commit()
            db.refresh(tx)
            assert db.get(SyncChange, tx.source_change_id).payload_json == dict(original, categoryId=tx.category_sync_id)
            assert preview_repair(db, ledger)['repair_count'] == 0
            pull = client.get('/api/v1/sync/pull', headers=h, params={'since': plan['base_change_id']})
            assert pull.status_code == 200 and tx.category_sync_id in pull.text
        assert tx.category_sync_id
        cats = client.get('/api/v1/read/workspace/categories', headers=h).json()
        cat = next(c for c in cats if c['id'] == tx.category_sync_id)
        assert cat['kind'] == 'income' and cat['tx_count'] == 1
        assert next(c for c in cats if c['kind'] == 'expense')['tx_count'] == 0
        rows = client.get('/api/v1/read/workspace/transactions', headers=h,
                          params={'category_sync_id': tx.category_sync_id}).json()
        assert rows['total'] == 1 and rows['items'][0]['tx_type'] == 'income'
        full = client.get('/api/v1/sync/full', headers=h, params={'ledger_id': external})
        assert full.status_code == 200 and tx.category_sync_id in full.text
    finally:
        db.close()

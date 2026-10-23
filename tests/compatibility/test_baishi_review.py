"""跨文件核对的接口契约，全部使用合成账单和内存数据库。"""

import io
import json

from openpyxl import load_workbook
from sqlalchemy import select
from test_baishi_aa import client as _client
from test_baishi_aa import upload, workbook
from test_import_csv import _login, _make_ledger

client = _client  # 复用隔离数据库夹具，避免重复维护测试环境。


def changed(payload, note=None, extra=False, book=None):
    w = load_workbook(io.BytesIO(payload))
    s = w.active
    s.insert_rows(1)  # 原始行号变化，不改变账单内容。
    if note is not None:
        s.cell(4, 17, note)
    if book:
        s.cell(4, 2, book)
    if extra:
        s.append([c.value for c in s[4]])
    out = io.BytesIO()
    w.save(out)
    return out.getvalue()


def review(c, token, key, **body):
    return c.post(
        f"/api/v1/import/{key}/review", json=body, headers={"Authorization": f"Bearer {token}"}
    )


def confirm_execute(c, token, summary, decisions=None):
    key = summary["import_token"]
    plan = review(c, token, key).json()
    r = review(c, token, key, revision=plan["revision"], decisions=decisions or {}, confirm=True)
    assert r.status_code == 200, r.text
    r = c.post(f"/api/v1/import/{key}/execute", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200, r.text
    blocks = [x for x in r.text.split("\n\n") if x.startswith("event: complete")]
    assert blocks, r.text
    return json.loads(blocks[0].split("data: ", 1)[1])


def test_execute_requires_explicit_review(client):
    token = _login(client)
    ledger = _make_ledger(client, token)
    summary = upload(client, token, workbook(), ledger)
    r = client.post(
        f"/api/v1/import/{summary['import_token']}/execute",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert r.status_code == 409
    assert r.json()["error_code"] == "IMPORT_REVIEW_REQUIRED"


def test_reexport_and_multiplicity(client):
    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook(duplicate=True)
    assert (
        confirm_execute(client, token, upload(client, token, payload, ledger))["created_tx_count"]
        == 2
    )
    second = upload(client, token, changed(payload), ledger)
    plan = review(client, token, second["import_token"]).json()
    assert plan["counts"]["matched"] == 2
    assert plan["counts"]["unresolved"] == 0
    assert confirm_execute(client, token, second)["created_tx_count"] == 0
    third = upload(client, token, changed(payload, extra=True), ledger)
    plan = review(client, token, third["import_token"]).json()
    assert plan["counts"]["matched"] == 2
    assert plan["counts"]["new"] == 1
    assert confirm_execute(client, token, third)["created_tx_count"] == 1


def test_changed_note_requires_choice_and_never_overwrites(client):
    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook()
    confirm_execute(client, token, upload(client, token, payload, ledger))
    summary = upload(client, token, changed(payload, note="不同备注"), ledger)
    plan = review(client, token, summary["import_token"]).json()
    assert plan["counts"]["review"] == 1
    row = plan["items"][0]
    r = review(client, token, summary["import_token"], revision=plan["revision"], confirm=True)
    assert r.status_code == 409
    assert (
        confirm_execute(client, token, summary, {str(row["row_number"]): {"action": "insert"}})[
            "created_tx_count"
        ]
        == 1
    )
    from src.database import get_db
    from src.main import app
    from src.models import ReadTxProjection

    with next(app.dependency_overrides[get_db]()) as db:
        assert len(db.scalars(select(ReadTxProjection)).all()) == 2


def test_skip_links_selected_record_for_future_reexports(client):
    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook()
    confirm_execute(client, token, upload(client, token, payload, ledger))
    summary = upload(client, token, changed(payload, note="修订备注"), ledger)
    plan = review(client, token, summary["import_token"]).json()
    row = plan["items"][0]
    decisions = {
        str(row["row_number"]): {"action": "skip", "match_id": row["candidates"][0]["sync_id"]}
    }
    assert confirm_execute(client, token, summary, decisions)["created_tx_count"] == 0
    again = upload(client, token, changed(changed(payload, note="修订备注")), ledger)
    assert review(client, token, again["import_token"]).json()["counts"]["matched"] == 1


def test_stale_confirmation_blocks_execution(client):
    token = _login(client)
    ledger = _make_ledger(client, token)
    summary = upload(client, token, workbook(), ledger)
    key = summary["import_token"]
    plan = review(client, token, key).json()
    assert review(client, token, key, revision=plan["revision"], confirm=True).status_code == 200
    written = client.post(
        f"/api/v1/write/ledgers/{ledger}/transactions",
        headers={"Authorization": f"Bearer {token}", "X-Device-ID": "d-web"},
        json={
            "base_change_id": 0,
            "tx_type": "expense",
            "amount": 3,
            "happened_at": "2026-01-03T00:00:00Z",
        },
    )
    assert written.status_code == 200, written.text
    r = client.post(f"/api/v1/import/{key}/execute", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 409
    assert r.json()["error_code"] == "IMPORT_REVIEW_STALE"


def test_review_owner_and_pagination(client):
    token = _login(client)
    other = _login(client, "other@example.com")
    ledger = _make_ledger(client, token)
    summary = upload(client, token, workbook(duplicate=True), ledger)
    key = summary["import_token"]
    assert review(client, other, key).status_code == 410
    plan = review(client, token, key, offset=1, limit=1).json()
    assert plan["total"] == 2
    assert len(plan["items"]) == 1


def test_deleted_target_is_not_treated_as_duplicate(client):
    from src.database import get_db
    from src.main import app
    from src.models import ReadTxProjection

    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook()
    confirm_execute(client, token, upload(client, token, payload, ledger))
    with next(app.dependency_overrides[get_db]()) as db:
        db.delete(db.scalar(select(ReadTxProjection)))
        db.commit()
    summary = upload(client, token, changed(payload), ledger)
    plan = review(client, token, summary["import_token"]).json()
    assert plan["counts"]["matched"] == 0
    assert plan["counts"]["new"] == 1


def test_manual_target_edits_are_preserved(client):
    from datetime import datetime, timezone

    from src.database import get_db
    from src.main import app
    from src.models import ReadTxProjection

    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook()
    confirm_execute(client, token, upload(client, token, payload, ledger))
    with next(app.dependency_overrides[get_db]()) as db:
        tx = db.scalar(select(ReadTxProjection))
        tx.note = "目标端人工修订"
        tx.happened_at = datetime(2026, 2, 1, tzinfo=timezone.utc)
        db.commit()
    summary = upload(client, token, changed(payload), ledger)
    assert review(client, token, summary["import_token"]).json()["counts"]["matched"] == 1
    assert confirm_execute(client, token, summary)["created_tx_count"] == 0
    with next(app.dependency_overrides[get_db]()) as db:
        tx = db.scalar(select(ReadTxProjection))
        assert tx.note == "目标端人工修订"
        assert tx.happened_at.month == 2


def test_unresolved_off_page_and_invalid_candidate_cannot_bypass(client):
    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook(duplicate=True)
    confirm_execute(client, token, upload(client, token, payload, ledger))
    summary = upload(client, token, changed(payload, note="修改"), ledger)
    key = summary["import_token"]
    plan = review(client, token, key, offset=1, limit=1).json()
    assert plan["counts"]["unresolved"] == 1
    assert plan["items"][0]["kind"] == "matched"
    assert review(client, token, key, revision=plan["revision"], confirm=True).status_code == 409
    result = review(
        client,
        token,
        key,
        revision=plan["revision"],
        decisions={"4": {"action": "skip", "match_id": "foreign-record"}},
    )
    assert result.status_code == 400
    assert review(client, token, key).json()["counts"]["unresolved"] == 1


def test_failed_import_rolls_back_source_links(client, monkeypatch):
    from src.database import get_db
    from src.main import app
    from src.models import ImportSourceRecord, ReadTxProjection
    from src.routers.import_data import endpoints

    token = _login(client)
    ledger = _make_ledger(client, token)
    summary = upload(client, token, workbook(), ledger)
    key = summary["import_token"]
    plan = review(client, token, key).json()
    assert review(client, token, key, revision=plan["revision"], confirm=True).status_code == 200

    def fail(*args, **kwargs):
        raise ValueError("synthetic rollback")

    monkeypatch.setattr(endpoints, "_emit_entity_diffs", fail)
    result = client.post(
        f"/api/v1/import/{key}/execute", headers={"Authorization": f"Bearer {token}"}
    )
    assert "event: error" in result.text
    with next(app.dependency_overrides[get_db]()) as db:
        assert not db.scalars(select(ImportSourceRecord)).all()
        assert not db.scalars(select(ReadTxProjection)).all()


def test_legacy_same_file_links_source_and_cross_file_needs_review(client):
    from src.database import get_db
    from src.main import app
    from src.models import ImportSourceRecord, ReadTxProjection
    from src.services.compatibility.baishi_aa import source_id
    from src.services.import_data.parser import parse_excel_bytes

    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook()
    confirm_execute(client, token, upload(client, token, payload, ledger))
    data = parse_excel_bytes(payload=payload)
    with next(app.dependency_overrides[get_db]()) as db:
        tx = db.scalar(select(ReadTxProjection))
        tx.sync_id = source_id(data.source_file_hash, 3)
        for record in db.scalars(select(ImportSourceRecord)):
            db.delete(record)
        db.commit()
    different = upload(client, token, changed(payload), ledger)
    assert review(client, token, different["import_token"]).json()["counts"]["review"] == 1
    same = upload(client, token, payload, ledger)
    plan = review(client, token, same["import_token"]).json()
    assert plan["items"][0]["reason"] == "same_file"
    assert confirm_execute(client, token, same)["created_tx_count"] == 0
    again = upload(client, token, changed(payload), ledger)
    assert review(client, token, again["import_token"]).json()["counts"]["matched"] == 1


def test_change_after_confirmation_before_write_lock_is_rejected(client, monkeypatch):
    from src.models import ReadTxProjection
    from src.routers.import_data import endpoints

    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook()
    confirm_execute(client, token, upload(client, token, payload, ledger))
    summary = upload(client, token, changed(payload), ledger)
    key = summary["import_token"]
    plan = review(client, token, key).json()
    assert review(client, token, key, revision=plan["revision"], confirm=True).status_code == 200
    original = endpoints.lock_ledger_for_materialize

    def lock_with_change(db, ledger_id):
        original(db, ledger_id)
        tx = db.scalar(select(ReadTxProjection))
        tx.note = "模拟确认后的写入"
        db.flush()

    monkeypatch.setattr(endpoints, "lock_ledger_for_materialize", lock_with_change)
    result = client.post(
        f"/api/v1/import/{key}/execute", headers={"Authorization": f"Bearer {token}"}
    )
    assert "IMPORT_REVIEW_STALE" in result.text
    assert "event: complete" not in result.text

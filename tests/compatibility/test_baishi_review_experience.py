"""核对分页、金额对账和历史权限，全部使用合成账单。"""

import io

from openpyxl import load_workbook
from sqlalchemy import select
from test_baishi_aa import client as _client
from test_baishi_aa import upload, workbook
from test_baishi_review import changed, confirm_execute, review
from test_import_csv import _login, _make_ledger

client = _client


def many_rows(count=27):
    w = load_workbook(io.BytesIO(workbook()))
    s = w.active
    row = [c.value for c in s[3]]
    for _ in range(count - 1):
        s.append(row)
    out = io.BytesIO()
    w.save(out)
    return out.getvalue()


def history(c, token, **params):
    return c.get(
        "/api/v1/import/history", headers={"Authorization": f"Bearer {token}"}, params=params
    )


def test_candidates_beyond_twenty_search_and_selection(client):
    token = _login(client)
    ledger = _make_ledger(client, token)
    original = many_rows()
    confirm_execute(client, token, upload(client, token, original, ledger))
    w = load_workbook(io.BytesIO(changed(original)))
    for row in w.active.iter_rows(min_row=4):
        row[16].value = "来源修订"
    out = io.BytesIO()
    w.save(out)
    summary = upload(client, token, out.getvalue(), ledger)
    key = summary["import_token"]
    plan = review(client, token, key).json()
    row = plan["items"][0]
    assert row["candidate_count"] == 27 and len(row["candidates"]) == 20
    page = review(client, token, key, candidate_queries={"4": {"offset": 20}}).json()["items"][0]
    assert len(page["candidates"]) == 7
    selected = page["candidates"][-1]["sync_id"]
    filtered = review(client, token, key, candidate_queries={"4": {"search": selected}}).json()[
        "items"
    ][0]
    assert filtered["candidate_total"] == 1
    assert filtered["candidates"][0]["differences"] == ["note"]
    decided = review(
        client,
        token,
        key,
        revision=plan["revision"],
        decisions={
            "4": {"action": "skip", "match_id": selected},
        },
    ).json()
    assert decided["items"][0]["selected_candidate"]["sync_id"] == selected
    assert decided["amounts"]["file"] == {"count": 27, "income": "0", "expense": "332.10"}
    assert decided["amounts"]["skip"]["expense"] == "12.30"
    assert decided["amounts"]["unresolved"]["expense"] == "319.80"
    assert decided["amounts"]["insert"]["expense"] == "0"
    assert "_live" not in decided
    assert not any(k.startswith("_") for k in decided["items"][0])


def test_unresolved_filter_keeps_global_counts_and_confirmation(client):
    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook(duplicate=True)
    confirm_execute(client, token, upload(client, token, payload, ledger))
    key = upload(client, token, changed(payload, note="待核对"), ledger)["import_token"]
    plan = review(client, token, key, unresolved_only=True, offset=999).json()
    assert plan["total"] == 1 and plan["file_total"] == 2
    assert plan["offset"] == 0 and plan["items"][0]["action"] is None
    assert plan["amounts"]["skip"]["expense"] == "12.30"
    assert plan["amounts"]["unresolved"]["expense"] == "12.30"
    assert (
        review(
            client, token, key, revision=plan["revision"], unresolved_only=True, confirm=True
        ).status_code
        == 409
    )
    decided = review(
        client,
        token,
        key,
        revision=plan["revision"],
        unresolved_only=True,
        decisions={"4": {"action": "insert"}},
    ).json()
    assert decided["total"] == 0 and decided["items"] == []
    assert decided["counts"]["insert"] == 1 and decided["counts"]["skip"] == 1
    assert decided["amounts"]["file"]["expense"] == "24.60"
    assert decided["amounts"]["insert"]["expense"] == "12.30"
    assert review(client, token, key, revision=decided["revision"], confirm=True).status_code == 200


def test_reconciliation_keeps_income_and_expense_separate_and_exact(client):
    token = _login(client)
    ledger = _make_ledger(client, token)
    w = load_workbook(io.BytesIO(workbook()))
    s = w.active
    s.cell(3, 7, "0.10")
    s.cell(3, 8, "0.10")
    s.cell(3, 12, "-0.10")
    s.cell(3, 14, "0.10")
    row = [c.value for c in s[3]]
    row[0], row[6], row[7], row[11], row[13] = "收入", "0.20", "0.20", "0.20", "0.20"
    row[3], row[4] = "收入", ""
    s.append(row)
    out = io.BytesIO()
    w.save(out)
    summary = upload(client, token, out.getvalue(), ledger)
    plan = review(client, token, summary["import_token"]).json()
    assert plan["amounts"]["insert"] == {"count": 2, "income": "0.20", "expense": "0.10"}
    assert plan["target"]["ledger_id"] == ledger and plan["currency"] == "CNY"


def test_history_reports_success_pagination_and_isolation(client):
    token = _login(client)
    other = _login(client, "other-history@example.com")
    ledger = _make_ledger(client, token)
    payload = workbook()
    confirm_execute(client, token, upload(client, token, payload, ledger))
    confirm_execute(client, token, upload(client, token, changed(payload), ledger))
    result = history(client, token, limit=1).json()
    assert result["total"] == 2 and len(result["items"]) == 1
    latest = result["items"][0]
    assert latest["result"] == "success" and latest["skipped_count"] == 1
    assert latest["report"]["amounts"]["skip"]["expense"] == "12.30"
    assert latest["report"]["amounts"]["insert"]["expense"] == "0"
    assert len(latest["file_hash"]) == 64
    assert history(client, token, offset=1, limit=1).json()["items"][0]["created_count"] == 1
    assert history(client, other).json()["items"] == []


def test_failed_execution_has_no_success_history(client, monkeypatch):
    from src.routers.import_data import endpoints

    token = _login(client)
    ledger = _make_ledger(client, token)
    summary = upload(client, token, workbook(), ledger)
    key = summary["import_token"]
    plan = review(client, token, key).json()
    review(client, token, key, revision=plan["revision"], confirm=True)

    def fail(*args, **kwargs):
        raise ValueError("synthetic failure")

    monkeypatch.setattr(endpoints, "_emit_entity_diffs", fail)
    response = client.post(
        f"/api/v1/import/{key}/execute", headers={"Authorization": f"Bearer {token}"}
    )
    assert "event: error" in response.text
    assert history(client, token).json()["total"] == 0


def test_history_hidden_after_member_permission_revoked(client):
    from src.database import get_db
    from src.main import app
    from src.models import Ledger, LedgerMember, User

    owner = _login(client)
    editor = _login(client, "editor-history@example.com")
    ledger = _make_ledger(client, owner)
    with next(app.dependency_overrides[get_db]()) as db:
        ledger_row = db.scalar(select(Ledger).where(Ledger.external_id == ledger))
        user = db.scalar(select(User).where(User.email == "editor-history@example.com"))
        member = LedgerMember(ledger_id=ledger_row.id, user_id=user.id, role="editor")
        db.add(member)
        db.commit()
    confirm_execute(client, editor, upload(client, editor, workbook(), ledger))
    assert history(client, editor).json()["total"] == 1
    assert history(client, owner).json()["total"] == 0
    with next(app.dependency_overrides[get_db]()) as db:
        member = db.scalar(select(LedgerMember).where(LedgerMember.role == "editor"))
        member.role = "viewer"
        db.commit()
    assert history(client, editor).json()["total"] == 0

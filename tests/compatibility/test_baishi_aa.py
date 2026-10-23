"""百事AA兼容：合成夹具，不包含私人账单。"""

import io
from decimal import Decimal

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from PIL import Image as PILImage

from src.services.import_data.parser import parse_excel_bytes
from src.services.import_data.transformer import apply_mapping

HEADERS = [
    "交易类型",
    "账本名称",
    "日期",
    "类别",
    "标签",
    "币种",
    "金额",
    "默认币种金额",
    "是否报销",
    "创建者",
    "付款人/收款人/交款人",
    "金额",
    "分账人/保管人",
    "金额",
    "账户",
    "参与统计标志",
    "描述",
]


def workbook(*, split="甲", images=False, duplicate=False):
    w = Workbook()
    s = w.active
    s.title = "合成账本"
    s.append(["百事AA记账 导出文件"] * 17)
    s.append(HEADERS)
    row = [
        "支出",
        "测试账本",
        "2026-01-02 10:30",
        "交通",
        "公交",
        "CNY",
        "12.30",
        "12.30",
        "",
        "甲",
        "甲",
        "-12.30",
        split,
        "12.30",
        "",
        "",
        "测试备注",
    ]
    s.append(row)
    if duplicate:
        s.append(row)
    if images:
        b = io.BytesIO()
        PILImage.new("RGB", (3, 3), "red").save(b, format="PNG")
        b.seek(0)
        s.add_image(Image(b), "R3")
    out = io.BytesIO()
    w.save(out)
    return out.getvalue()


def test_detect_duplicate_columns_and_mapping():
    data = parse_excel_bytes(payload=workbook())
    assert data.source_format == "baishi_aa"
    assert len(data.rows) == 1
    data.suggested_mapping.tz_offset_minutes = 480
    txs, errors, _ = apply_mapping(rows=data.rows, mapping=data.suggested_mapping)
    assert not errors
    assert txs[0].amount == Decimal("12.30")
    assert txs[0].parent_category_name == "交通"
    assert txs[0].category_name == "公交"
    assert txs[0].happened_at.hour == 2
    assert data.source_ledger_name == "测试账本"


def test_picture_is_bound_to_source_row():
    data = parse_excel_bytes(payload=workbook(images=True))
    assert len(data.attachments) == 1
    assert data.attachments[0].row_number == 3
    assert data.attachments[0].data


def test_unsupported_split_blocks_import():
    data = parse_excel_bytes(payload=workbook(split="乙"))
    assert any(e.code == "BAISHI_UNSUPPORTED" for e in data.parse_errors)


def test_legitimate_same_amount_time_rows_preserved():
    data = parse_excel_bytes(payload=workbook(duplicate=True))
    assert len(data.rows) == 2
    assert data.rows[0].row_number != data.rows[1].row_number


@pytest.fixture
def client():
    from test_import_csv import _make_client

    from src.main import app
    from src.services.import_data.cache import clear_all

    c = _make_client()
    clear_all()
    yield c
    c.close()
    app.dependency_overrides.clear()
    clear_all()


def upload(client, token, payload, ledger=None):
    r = client.post(
        "/api/v1/import/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "aa.xlsx",
                payload,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"target_ledger_id": ledger} if ledger else {},
    )
    assert r.status_code == 200, r.text
    return r.json()


def execute(client, token, key):
    r = client.post(f"/api/v1/import/{key}/execute", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200, r.text
    return r.text


def test_existing_import_repeat_and_legitimate_duplicates(client):
    from test_import_csv import _login, _make_ledger

    token = _login(client)
    ledger = _make_ledger(client, token)
    payload = workbook(duplicate=True)
    summary = upload(client, token, payload, ledger)
    assert summary["source_format"] == "baishi_aa"
    assert summary["stats"]["parse_errors_total"] == 0
    result = execute(client, token, summary["import_token"])
    assert '"created_tx_count": 2' in result, result
    summary = upload(client, token, payload, ledger)
    assert summary["stats"]["skipped_dedup"] == 2
    result = execute(client, token, summary["import_token"])
    assert '"created_tx_count": 0' in result, result


def test_new_ledger_with_picture_and_user_isolation(client, tmp_path, monkeypatch):
    from test_import_csv import _login

    from src.config import get_settings

    monkeypatch.setattr(get_settings(), "attachment_storage_dir", str(tmp_path / "attachments"))
    token = _login(client)
    other = _login(client, "other@test.com")
    summary = upload(client, token, workbook(images=True))
    key = summary["import_token"]
    r = client.post(
        f"/api/v1/import/{key}/preview",
        json={"new_ledger_name": "迁移测试"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["new_ledger_name"] == "迁移测试"
    assert r.json()["attachment_count"] == 1
    assert (
        client.post(
            f"/api/v1/import/{key}/execute", headers={"Authorization": f"Bearer {other}"}
        ).status_code
        == 410
    )
    result = execute(client, token, key)
    assert '"created_tx_count": 1' in result, result
    assert len(list((tmp_path / "attachments").rglob("*.png"))) == 1


def test_failure_rolls_back_new_ledger_and_files(client, tmp_path, monkeypatch):
    from sqlalchemy import select
    from test_import_csv import _login

    from src.config import get_settings
    from src.database import get_db
    from src.main import app
    from src.models import AttachmentFile, Ledger
    from src.routers.import_data import endpoints

    monkeypatch.setattr(get_settings(), "attachment_storage_dir", str(tmp_path / "attachments"))
    token = _login(client)
    summary = upload(client, token, workbook(images=True))
    key = summary["import_token"]
    client.post(
        f"/api/v1/import/{key}/preview",
        json={"new_ledger_name": "不应残留"},
        headers={"Authorization": f"Bearer {token}"},
    )

    def fail(*args, **kwargs):
        raise ValueError("synthetic failure")

    monkeypatch.setattr(endpoints, "_emit_entity_diffs", fail)
    result = execute(client, token, key)
    assert "event: error" in result, result
    with next(app.dependency_overrides[get_db]()) as db:
        assert not db.scalars(select(Ledger)).all()
        assert not db.scalars(select(AttachmentFile)).all()
    assert not list((tmp_path / "attachments").rglob("*.png"))


def test_zero_amount_is_preserved():
    w = Workbook()
    s = w.active
    s.append(HEADERS)
    s.append(
        [
            "支出",
            "测试",
            "2026-01-02 10:30",
            "其他",
            "",
            "CNY",
            "0",
            "0",
            "",
            "甲",
            "甲",
            "0",
            "甲",
            "0",
            "",
            "",
            "免费项目",
        ]
    )
    out = io.BytesIO()
    w.save(out)
    data = parse_excel_bytes(payload=out.getvalue())
    assert not data.parse_errors
    txs, errors, _ = apply_mapping(rows=data.rows, mapping=data.suggested_mapping)
    assert not errors and txs[0].amount == 0


def test_category_name_collision_is_explicitly_qualified():
    w = Workbook()
    s = w.active
    s.append(HEADERS)
    for cat, sub in [("交通", "其他"), ("其他", "")]:
        s.append(
            [
                "支出",
                "测试",
                "2026-01-02 10:30",
                cat,
                sub,
                "CNY",
                "1",
                "1",
                "",
                "甲",
                "甲",
                "-1",
                "甲",
                "1",
                "",
                "",
                "",
            ]
        )
    out = io.BytesIO()
    w.save(out)
    data = parse_excel_bytes(payload=out.getvalue())
    assert data.rows[0].cells["二级分类"] == "其他（交通）"
    assert any(w.code == "BAISHI_CATEGORY_QUALIFIED" for w in data.parse_warnings)


def test_corrupt_xlsx_returns_parse_error(client):
    from test_import_csv import _login

    token = _login(client)
    r = client.post(
        "/api/v1/import/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("bad.xlsx", b"not-a-zip", "application/octet-stream")},
    )
    assert r.status_code == 400
    assert r.json()["error_code"] == "IMPORT_XLSX_PARSE_FAILED"


def test_unsupported_source_cannot_create_ledger(client):
    from sqlalchemy import select
    from test_import_csv import _login

    from src.database import get_db
    from src.main import app
    from src.models import Ledger

    token = _login(client)
    summary = upload(client, token, workbook(split="乙"))
    key = summary["import_token"]
    client.post(
        f"/api/v1/import/{key}/preview",
        json={"new_ledger_name": "禁止写入"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert "event: error" in execute(client, token, key)
    with next(app.dependency_overrides[get_db]()) as db:
        assert not db.scalars(select(Ledger)).all()


def test_busy_token_rejects_execution_and_preview(client):
    from test_import_csv import _login, _make_ledger

    from src.services.import_data.cache import cancel_token, claim_execution, release_execution

    token = _login(client)
    ledger = _make_ledger(client, token)
    key = upload(client, token, workbook(), ledger)["import_token"]
    from sqlalchemy import select

    from src.database import get_db
    from src.main import app
    from src.models import User

    with next(app.dependency_overrides[get_db]()) as db:
        user_id = db.scalar(select(User)).id
    assert claim_execution(token=key, user_id=user_id)
    assert not claim_execution(token=key, user_id=user_id)
    assert not cancel_token(token=key, user_id=user_id)
    headers = {"Authorization": f"Bearer {token}"}
    assert client.post(f"/api/v1/import/{key}/execute", headers=headers).status_code == 409
    assert client.post(f"/api/v1/import/{key}/preview", headers=headers, json={}).status_code == 409
    release_execution(token=key, user_id=user_id)
    assert "event: complete" in execute(client, token, key)
    assert client.post(f"/api/v1/import/{key}/execute", headers=headers).status_code == 410


def test_target_currency_and_category_conflicts():
    from src.services.compatibility.baishi_aa import validate_target

    data = parse_excel_bytes(payload=workbook())
    txs, _, _ = apply_mapping(rows=data.rows, mapping=data.suggested_mapping)
    assert validate_target(txs, {}, "USD")
    snapshot = {"categories": [{"name": "公交", "kind": "expense", "parentName": "其他"}]}
    assert any(e.code == "BAISHI_CATEGORY_CONFLICT" for e in validate_target(txs, snapshot, "CNY"))


def test_viewer_and_other_user_cannot_import(client):
    from sqlalchemy import select
    from test_import_csv import _login, _make_ledger

    from src.database import get_db
    from src.main import app
    from src.models import Ledger, LedgerMember, User

    owner = _login(client)
    viewer = _login(client, "viewer@test.com")
    outsider = _login(client, "outsider@test.com")
    ledger = _make_ledger(client, owner)
    with next(app.dependency_overrides[get_db]()) as db:
        ledger_row = db.scalar(select(Ledger).where(Ledger.external_id == ledger))
        viewer_row = db.scalar(select(User).where(User.email == "viewer@test.com"))
        db.add(LedgerMember(ledger_id=ledger_row.id, user_id=viewer_row.id, role="viewer"))
        db.commit()
    for token in (viewer, outsider):
        r = client.post(
            "/api/v1/import/upload",
            headers={"Authorization": f"Bearer {token}"},
            files={"file": ("aa.xlsx", workbook(), "application/octet-stream")},
            data={"target_ledger_id": ledger},
        )
        assert r.status_code == 404
